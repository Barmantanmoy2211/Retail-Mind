"""RetailFlow AI - Backend integration tests.

Covers: auth, platform admin, business CRUD, POS billing, inventory,
purchase orders, expenses, dashboard, RBAC, tenant isolation.
"""
import os
import time
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://retailflow-dashboard.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

CREDS = {
    "platform_admin": ("admin@retailflow.ai", "Admin@123"),
    "business_admin": ("owner@bharatmart.com", "Owner@123"),
    "outlet_manager": ("manager@bharatmart.com", "Manager@123"),
    "cashier": ("cashier@bharatmart.com", "Cashier@123"),
}

# ---------- fixtures ----------
@pytest.fixture(scope="session")
def session():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


def _login(s, role):
    email, pw = CREDS[role]
    r = s.post(f"{API}/auth/login", json={"email": email, "password": pw}, timeout=15)
    assert r.status_code == 200, f"Login failed for {role}: {r.text}"
    return r.json()


@pytest.fixture(scope="session")
def tokens(session):
    out = {}
    for role in CREDS:
        out[role] = _login(session, role)
    return out


def _client(token):
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json", "Authorization": f"Bearer {token}"})
    return s


# ---------- Health ----------
def test_health(session):
    r = session.get(f"{API}/health", timeout=10)
    assert r.status_code == 200
    assert r.json()["status"] == "ok"


# ---------- Auth ----------
@pytest.mark.parametrize("role", list(CREDS.keys()))
def test_login_all_roles(session, role):
    data = _login(session, role)
    assert "access_token" in data
    assert data["user"]["email"] == CREDS[role][0]
    assert data["user"]["role"] == role
    if role == "platform_admin":
        assert data["business"] is None
    else:
        assert data["business"] is not None
        assert data["business"]["subscription_status"] == "approved"


def test_login_invalid(session):
    r = session.post(f"{API}/auth/login", json={"email": "owner@bharatmart.com", "password": "wrong"})
    assert r.status_code == 401


def test_auth_me(tokens):
    c = _client(tokens["business_admin"]["access_token"])
    r = c.get(f"{API}/auth/me")
    assert r.status_code == 200
    assert r.json()["user"]["email"] == CREDS["business_admin"][0]


def test_register_business(session):
    ts = int(time.time())
    payload = {
        "business_name": f"TEST_Biz_{ts}",
        "business_type": "Grocery",
        "owner_name": "Test Owner",
        "email": f"test_owner_{ts}@example.com",
        "phone": "9999999999",
        "password": "Test@123",
        "plan": "starter",
    }
    r = session.post(f"{API}/auth/register-business", json=payload)
    assert r.status_code == 200
    assert r.json()["success"] is True
    # Pending business should fail to login? Actually current code allows login for pending
    r2 = session.post(f"{API}/auth/login", json={"email": payload["email"], "password": "Test@123"})
    assert r2.status_code == 200
    assert r2.json()["business"]["subscription_status"] == "pending"


# ---------- Platform admin ----------
def test_platform_stats(tokens):
    c = _client(tokens["platform_admin"]["access_token"])
    r = c.get(f"{API}/platform/stats")
    assert r.status_code == 200
    d = r.json()
    for key in ["total_businesses", "active_businesses", "pending_approvals",
                "total_outlets", "total_customers", "total_bills",
                "total_revenue", "plan_distribution", "top_businesses",
                "recent_registrations"]:
        assert key in d, f"missing {key}"


def test_platform_list_businesses(tokens):
    c = _client(tokens["platform_admin"]["access_token"])
    r = c.get(f"{API}/platform/businesses")
    assert r.status_code == 200
    bizs = r.json()
    assert len(bizs) >= 1
    assert "outlets_count" in bizs[0]
    assert "users_count" in bizs[0]


def test_platform_approve_pending(session, tokens):
    # Create a pending biz then approve
    ts = int(time.time())
    payload = {
        "business_name": f"TEST_Approve_{ts}",
        "business_type": "Pharmacy",
        "owner_name": "Approve Test",
        "email": f"approve_{ts}@example.com",
        "phone": "8888888888",
        "password": "Test@123",
        "plan": "starter",
    }
    session.post(f"{API}/auth/register-business", json=payload)
    c = _client(tokens["platform_admin"]["access_token"])
    bizs = c.get(f"{API}/platform/businesses", params={"status": "pending"}).json()
    target = next((b for b in bizs if b["email"] == payload["email"]), None)
    assert target is not None
    r = c.post(f"{API}/platform/businesses/{target['id']}/action", json={"action": "approve"})
    assert r.status_code == 200
    # Verify
    bizs2 = c.get(f"{API}/platform/businesses").json()
    updated = next(b for b in bizs2 if b["id"] == target["id"])
    assert updated["subscription_status"] == "approved"


# ---------- Outlets ----------
def test_list_outlets_business_admin(tokens):
    c = _client(tokens["business_admin"]["access_token"])
    r = c.get(f"{API}/outlets")
    assert r.status_code == 200
    outlets = r.json()
    assert len(outlets) >= 1


def test_cashier_cannot_create_outlet(tokens):
    c = _client(tokens["cashier"]["access_token"])
    r = c.post(f"{API}/outlets", json={"name": "TEST_x", "address": "x"})
    assert r.status_code == 403


def test_outlet_manager_sees_only_own_outlet(tokens):
    c = _client(tokens["outlet_manager"]["access_token"])
    r = c.get(f"{API}/outlets")
    assert r.status_code == 200
    outlets = r.json()
    assert len(outlets) == 1  # filtered to assigned outlet


# ---------- Products CRUD ----------
def test_product_crud(tokens):
    c = _client(tokens["business_admin"]["access_token"])
    ts = int(time.time())
    payload = {
        "name": f"TEST_Product_{ts}", "sku": f"TSKU{ts}", "selling_price": 100.0,
        "cost_price": 60.0, "tax_percent": 5.0, "min_stock": 5, "unit": "pcs",
    }
    r = c.post(f"{API}/products", json=payload)
    assert r.status_code == 200
    pid = r.json()["id"]

    r2 = c.get(f"{API}/products")
    assert r2.status_code == 200
    assert any(p["id"] == pid for p in r2.json())

    r3 = c.put(f"{API}/products/{pid}", json={"selling_price": 120.0})
    assert r3.status_code == 200

    r4 = c.delete(f"{API}/products/{pid}")
    assert r4.status_code == 200


# ---------- Customers ----------
def test_customer_create_and_get(tokens):
    c = _client(tokens["business_admin"]["access_token"])
    ts = int(time.time())
    r = c.post(f"{API}/customers", json={"name": f"TEST_Cust_{ts}", "phone": f"7{ts}"})
    assert r.status_code == 200
    cid = r.json()["id"]
    r2 = c.get(f"{API}/customers/{cid}")
    assert r2.status_code == 200
    assert r2.json()["phone"] == f"7{ts}"


# ---------- POS / Bill ----------
def test_create_bill_full_flow(tokens):
    c = _client(tokens["business_admin"]["access_token"])
    outlets = c.get(f"{API}/outlets").json()
    outlet_id = outlets[0]["id"]
    products = c.get(f"{API}/products", params={"outlet_id": outlet_id}).json()
    if not products:
        products = c.get(f"{API}/products").json()
    assert len(products) >= 1
    prod = products[0]
    pid = prod["id"]

    # Ensure stock for that product+outlet
    c.post(f"{API}/inventory/transactions", json={
        "product_id": pid, "outlet_id": outlet_id, "type": "stock_in", "quantity": 20
    })

    payload = {
        "outlet_id": outlet_id,
        "items": [{"product_id": pid, "quantity": 2, "discount": 0}],
        "discount_total": 0,
        "payment_method": "cash",
    }
    r = c.post(f"{API}/bills", json=payload)
    assert r.status_code == 200, r.text
    bill = r.json()
    assert "bill_no" in bill and bill["bill_no"].startswith("INV-")
    assert bill["total"] > 0
    expected_subtotal = prod["selling_price"] * 2
    assert abs(bill["subtotal"] - expected_subtotal) < 0.01

    # Verify retrieving
    r2 = c.get(f"{API}/bills/{bill['id']}")
    assert r2.status_code == 200


def test_cashier_cannot_bill_other_outlet(tokens):
    c_admin = _client(tokens["business_admin"]["access_token"])
    outlets = c_admin.get(f"{API}/outlets").json()
    cashier = tokens["cashier"]["user"]
    other_outlet = next((o for o in outlets if o["id"] != cashier.get("outlet_id")), None)
    if not other_outlet:
        pytest.skip("no other outlet")
    products = c_admin.get(f"{API}/products").json()
    if not products:
        pytest.skip("no products")
    c = _client(tokens["cashier"]["access_token"])
    r = c.post(f"{API}/bills", json={
        "outlet_id": other_outlet["id"],
        "items": [{"product_id": products[0]["id"], "quantity": 1}],
        "payment_method": "cash",
    })
    assert r.status_code == 403


# ---------- Inventory ----------
def test_inventory_stock_in_updates_stock(tokens):
    c = _client(tokens["business_admin"]["access_token"])
    outlets = c.get(f"{API}/outlets").json()
    outlet_id = outlets[0]["id"]
    products = c.get(f"{API}/products").json()
    pid = products[0]["id"]
    before = products[0].get("stock_by_outlet", {}).get(outlet_id, 0)
    r = c.post(f"{API}/inventory/transactions", json={
        "product_id": pid, "outlet_id": outlet_id, "type": "stock_in", "quantity": 5
    })
    assert r.status_code == 200
    products2 = c.get(f"{API}/products").json()
    after = next(p for p in products2 if p["id"] == pid)["stock_by_outlet"].get(outlet_id, 0)
    assert after == before + 5


def test_low_stock_endpoint(tokens):
    c = _client(tokens["business_admin"]["access_token"])
    r = c.get(f"{API}/inventory/low-stock")
    assert r.status_code == 200
    assert isinstance(r.json(), list)


# ---------- Suppliers + PO ----------
def test_supplier_and_po_received_creates_expense(tokens):
    c = _client(tokens["business_admin"]["access_token"])
    ts = int(time.time())
    r = c.post(f"{API}/suppliers", json={"name": f"TEST_Supp_{ts}", "contact": "9000000"})
    assert r.status_code == 200
    sid = r.json()["id"]

    outlets = c.get(f"{API}/outlets").json()
    outlet_id = outlets[0]["id"]
    products = c.get(f"{API}/products").json()
    pid = products[0]["id"]
    before = products[0].get("stock_by_outlet", {}).get(outlet_id, 0)

    r2 = c.post(f"{API}/purchase-orders", json={
        "supplier_id": sid, "outlet_id": outlet_id,
        "items": [{"product_id": pid, "quantity": 10, "cost_price": 50.0}],
    })
    assert r2.status_code == 200
    po_id = r2.json()["id"]
    r3 = c.put(f"{API}/purchase-orders/{po_id}", json={"status": "received"})
    assert r3.status_code == 200

    products2 = c.get(f"{API}/products").json()
    after = next(p for p in products2 if p["id"] == pid)["stock_by_outlet"].get(outlet_id, 0)
    assert after == before + 10

    exps = c.get(f"{API}/expenses").json()
    assert any(po_id in (e.get("description") or "") for e in exps)


# ---------- Expenses ----------
def test_expense_admin_auto_approved(tokens):
    c = _client(tokens["business_admin"]["access_token"])
    outlets = c.get(f"{API}/outlets").json()
    r = c.post(f"{API}/expenses", json={
        "outlet_id": outlets[0]["id"], "category": "operations",
        "amount": 100.0, "description": "TEST expense admin",
    })
    assert r.status_code == 200
    exps = c.get(f"{API}/expenses").json()
    eid = r.json()["id"]
    e = next(e for e in exps if e["id"] == eid)
    assert e["status"] == "approved"


def test_expense_manager_pending_then_approve(tokens):
    cm = _client(tokens["outlet_manager"]["access_token"])
    outlet_id = tokens["outlet_manager"]["user"]["outlet_id"]
    r = cm.post(f"{API}/expenses", json={
        "outlet_id": outlet_id, "category": "operations",
        "amount": 50.0, "description": "TEST expense manager",
    })
    assert r.status_code == 200, r.text
    eid = r.json()["id"]
    ca = _client(tokens["business_admin"]["access_token"])
    r2 = ca.put(f"{API}/expenses/{eid}/action", json={"action": "approve"})
    assert r2.status_code == 200


# ---------- Dashboard ----------
def test_business_dashboard_kpis(tokens):
    c = _client(tokens["business_admin"]["access_token"])
    r = c.get(f"{API}/dashboard/business")
    assert r.status_code == 200
    d = r.json()
    for k in ["today_sales", "monthly_sales", "daily_trend",
              "outlet_comparison", "recent_bills", "top_products"]:
        assert k in d, f"missing dashboard key: {k}"


# ---------- Reports ----------
@pytest.mark.parametrize("rpath", ["profit-loss", "tax", "inventory", "sales"])
def test_reports_endpoints(tokens, rpath):
    c = _client(tokens["business_admin"]["access_token"])
    r = c.get(f"{API}/reports/{rpath}")
    assert r.status_code == 200, f"{rpath}: {r.text}"


# ---------- RBAC ----------
def test_cashier_cannot_create_user(tokens):
    c = _client(tokens["cashier"]["access_token"])
    r = c.post(f"{API}/users", json={
        "name": "x", "email": "x@x.com", "password": "x", "role": "cashier"
    })
    assert r.status_code == 403


def test_cashier_cannot_list_users(tokens):
    c = _client(tokens["cashier"]["access_token"])
    r = c.get(f"{API}/users")
    assert r.status_code == 403


def test_no_token_unauthorized(session):
    r = session.get(f"{API}/auth/me")
    assert r.status_code == 401


# ---------- Tenant isolation ----------
def test_tenant_isolation_products(session, tokens):
    """Create a separate biz, approve it, login, check it cannot see Bharat Mart products."""
    ts = int(time.time())
    payload = {
        "business_name": f"TEST_IsoBiz_{ts}",
        "business_type": "Retail",
        "owner_name": "Iso Owner",
        "email": f"iso_{ts}@example.com",
        "phone": "7000000000",
        "password": "Iso@1234",
        "plan": "starter",
    }
    session.post(f"{API}/auth/register-business", json=payload)
    pa = _client(tokens["platform_admin"]["access_token"])
    bizs = pa.get(f"{API}/platform/businesses", params={"status": "pending"}).json()
    target = next(b for b in bizs if b["email"] == payload["email"])
    pa.post(f"{API}/platform/businesses/{target['id']}/action", json={"action": "approve"})
    r = session.post(f"{API}/auth/login", json={"email": payload["email"], "password": "Iso@1234"})
    assert r.status_code == 200
    new_token = r.json()["access_token"]
    cnew = _client(new_token)
    products_new = cnew.get(f"{API}/products").json()
    cust_new = cnew.get(f"{API}/customers").json()
    # New biz should have no products/customers
    assert products_new == []
    assert cust_new == []


# ============================================================
# Phase 1.5 — Health Score, Excel Exports, S3 Upload skeleton
# ============================================================
import io
from openpyxl import load_workbook


# ---------- Health Score ----------
class TestHealthScore:
    def test_health_score_returns_outlets_with_metrics(self, tokens):
        c = _client(tokens["business_admin"]["access_token"])
        r = c.get(f"{API}/health-score", timeout=60)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "outlets" in data
        outlets = data["outlets"]
        assert len(outlets) >= 1, "Expected at least 1 outlet"
        # Sorted descending by score
        scores = [o["score"] for o in outlets]
        assert scores == sorted(scores, reverse=True), "Outlets must be sorted by score desc"

        required_metrics = {
            "sales_growth", "profit_margin", "customer_growth",
            "inventory_health", "expense_control", "tax_compliance",
            "reward_engagement",
        }
        valid_statuses = {"Excellent", "Good", "Needs Attention", "Critical"}

        for o in outlets:
            for k in ["outlet_id", "outlet_name", "score", "status",
                      "metrics", "recommendations", "revenue_30d", "profit_30d"]:
                assert k in o, f"missing key {k} in outlet result"
            assert o["status"] in valid_statuses, f"invalid status: {o['status']}"
            assert 0 <= o["score"] <= 100, f"score out of range: {o['score']}"
            assert set(o["metrics"].keys()) >= required_metrics, \
                f"missing metrics: {required_metrics - set(o['metrics'].keys())}"
            assert isinstance(o["recommendations"], list) and len(o["recommendations"]) >= 1

    def test_health_score_outlet_manager_scoped(self, tokens):
        c = _client(tokens["outlet_manager"]["access_token"])
        r = c.get(f"{API}/health-score", timeout=60)
        # Should succeed (manager has business_scope) - returns at least the assigned outlet
        assert r.status_code == 200, r.text
        assert "outlets" in r.json()

    def test_health_score_unauthorized(self, session):
        r = session.get(f"{API}/health-score")
        assert r.status_code == 401


# ---------- Excel Exports ----------
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _assert_xlsx(resp, expected_filename_part):
    assert resp.status_code == 200, resp.text[:300]
    assert XLSX_MEDIA in resp.headers.get("content-type", ""), \
        f"Wrong content-type: {resp.headers.get('content-type')}"
    cd = resp.headers.get("content-disposition", "")
    assert "attachment" in cd.lower(), f"Missing attachment in: {cd}"
    assert expected_filename_part in cd, f"Expected '{expected_filename_part}' in CD: {cd}"
    # Verify body is a valid XLSX (load via openpyxl)
    wb = load_workbook(io.BytesIO(resp.content))
    return wb


class TestExcelExports:
    def test_export_bills_xlsx(self, tokens):
        c = _client(tokens["business_admin"]["access_token"])
        # Use raw requests to keep bytes
        r = requests.get(
            f"{API}/exports/bills.xlsx",
            headers={"Authorization": f"Bearer {tokens['business_admin']['access_token']}"},
            timeout=60,
        )
        wb = _assert_xlsx(r, "bills-")
        assert "Bills" in wb.sheetnames
        ws = wb["Bills"]
        headers = [c.value for c in ws[1]]
        for h in ["Bill No", "Date", "Customer", "Total"]:
            assert h in headers, f"missing column {h}"

    def test_export_products_xlsx(self, tokens):
        r = requests.get(
            f"{API}/exports/products.xlsx",
            headers={"Authorization": f"Bearer {tokens['business_admin']['access_token']}"},
            timeout=60,
        )
        wb = _assert_xlsx(r, "products-")
        assert "Products" in wb.sheetnames
        # Demo data: 20 products → at least header + 1 row
        assert wb["Products"].max_row >= 2

    def test_export_customers_xlsx(self, tokens):
        r = requests.get(
            f"{API}/exports/customers.xlsx",
            headers={"Authorization": f"Bearer {tokens['business_admin']['access_token']}"},
            timeout=60,
        )
        wb = _assert_xlsx(r, "customers-")
        assert "Customers" in wb.sheetnames

    def test_export_profit_loss_xlsx_monthly(self, tokens):
        r = requests.get(
            f"{API}/exports/profit-loss.xlsx",
            params={"period": "monthly"},
            headers={"Authorization": f"Bearer {tokens['business_admin']['access_token']}"},
            timeout=60,
        )
        wb = _assert_xlsx(r, "profit-loss-monthly-")
        # Two sheets: P&L + Expenses by Category
        assert "P&L" in wb.sheetnames
        assert "Expenses by Category" in wb.sheetnames

    def test_export_unauthorized(self, session):
        r = session.get(f"{API}/exports/bills.xlsx")
        assert r.status_code == 401


# ---------- Uploads (S3 not configured branch) ----------
class TestUploads:
    def test_uploads_status_not_configured(self, tokens):
        c = _client(tokens["business_admin"]["access_token"])
        r = c.get(f"{API}/uploads/status")
        assert r.status_code == 200
        assert r.json() == {"s3_configured": False}, r.json()

    def test_upload_image_returns_503_when_s3_not_configured(self, tokens):
        # Tiny valid PNG (1x1)
        png_bytes = (
            b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
            b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\xf8\xcf"
            b"\xc0\x00\x00\x00\x03\x00\x01\xa0\x18\xe8\xd9\x00\x00\x00\x00IEND\xaeB`\x82"
        )
        r = requests.post(
            f"{API}/uploads/image",
            headers={"Authorization": f"Bearer {tokens['business_admin']['access_token']}"},
            files={"file": ("test.png", png_bytes, "image/png")},
            timeout=30,
        )
        assert r.status_code == 503, f"Expected 503, got {r.status_code}: {r.text}"
        body = r.json()
        detail = body.get("detail", "")
        assert "S3 not configured" in detail or "s3" in detail.lower(), \
            f"Detail does not mention S3: {detail}"

    def test_upload_image_rejects_non_image_content_type(self, tokens):
        r = requests.post(
            f"{API}/uploads/image",
            headers={"Authorization": f"Bearer {tokens['business_admin']['access_token']}"},
            files={"file": ("test.txt", b"hello world", "text/plain")},
            timeout=30,
        )
        assert r.status_code == 400, f"Expected 400, got {r.status_code}: {r.text}"

    def test_upload_image_rejects_oversize(self, tokens):
        # 6 MB of image-typed bytes
        big = b"\x89PNG\r\n\x1a\n" + (b"A" * (6 * 1024 * 1024))
        r = requests.post(
            f"{API}/uploads/image",
            headers={"Authorization": f"Bearer {tokens['business_admin']['access_token']}"},
            files={"file": ("big.png", big, "image/png")},
            timeout=60,
        )
        assert r.status_code == 400, f"Expected 400 for oversize, got {r.status_code}"
        assert "5MB" in r.text or "size" in r.text.lower()

    def test_upload_image_requires_auth(self, session):
        r = session.post(f"{API}/uploads/image")
        assert r.status_code == 401


# ---------- Organization / Permissions ----------
class TestOrganization:
    def test_auth_me_includes_permissions(self, tokens):
        c = _client(tokens["business_admin"]["access_token"])
        r = c.get(f"{API}/auth/me")
        assert r.status_code == 200
        body = r.json()
        assert "permissions" in body
        assert "*" in body["permissions"] or len(body["permissions"]) > 0

    def test_permissions_catalog(self, tokens):
        c = _client(tokens["business_admin"]["access_token"])
        r = c.get(f"{API}/org/permissions/catalog")
        assert r.status_code == 200
        assert "organization" in r.json()

    def test_list_roles_owner(self, tokens):
        c = _client(tokens["business_admin"]["access_token"])
        r = c.get(f"{API}/org/roles")
        assert r.status_code == 200
        assert isinstance(r.json(), list)

    def test_cashier_cannot_create_roles(self, tokens):
        c = _client(tokens["cashier"]["access_token"])
        r = c.post(f"{API}/org/roles", json={
            "name": "Test Role",
            "permissions": ["bills.create"],
            "scope": "outlet",
        })
        assert r.status_code == 403

    def test_org_chart(self, tokens):
        c = _client(tokens["business_admin"]["access_token"])
        r = c.get(f"{API}/org/chart")
        assert r.status_code == 200
        body = r.json()
        assert "employee_tree" in body
        assert "role_tree" in body

    def test_pending_approvals_endpoint(self, tokens):
        c = _client(tokens["business_admin"]["access_token"])
        r = c.get(f"{API}/org/approvals/pending")
        assert r.status_code == 200
        assert isinstance(r.json(), list)

